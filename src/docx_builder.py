"""Sprint 3 — docx_builder.py

Pure render layer: takes the approved logical POE (document_data + document_text)
and produces an institutional .docx using a provided template.

Design constraints:
  - No content validation.  No plan logic.  No API calls.
  - Preserves header, footer and institutional tables from the template by
    removing only <w:p> elements from the body, leaving <w:tbl> and <w:sectPr>.
  - Falls back to raw document_text when document_data has no structured sections.
  - Returns a result dict (never raises on missing optional sections).
  - Raises DocxBuildError on technical failures (missing template, unreadable
    file, write error).

Public API:
    DocxBuildError(Exception)
    build_docx_from_template(
        template_path: str,
        output_path: str,
        document_data: dict,
        document_text: str | None = None,
    ) -> dict

Return contract:
  {
    "ok":               bool,
    "output_path":      str,
    "sections_written": list[str],
    "warnings":         list[str],
  }
"""
import logging
from pathlib import Path

try:
    from docx import Document
    _DOCX_AVAILABLE = True
except ImportError:  # pragma: no cover
    _DOCX_AVAILABLE = False

logger = logging.getLogger(__name__)

# ---------------------------------------------------------------------------
# Section render order — canonical key → display heading label
# ---------------------------------------------------------------------------

_SECTION_RENDER_ORDER: list[tuple[str, str]] = [
    ("objetivo",               "1. Objetivo"),
    ("alcance",                "2. Alcance"),
    ("responsabilidades",      "3. Responsabilidades"),
    ("procedimiento",          "4. Procedimiento"),
    ("gestion_desviaciones",   "5. Gestión de Desviaciones y Retorno a Operación"),
    ("registros",              "6. Registros"),
    # Optional sections — rendered only if present in document_data
    ("definiciones",           "7. Definiciones"),
    ("referencias",            "8. Referencias"),
    ("control_cambios",        "9. Control de Cambios"),
    ("historial_modificaciones", "10. Historial de Modificaciones"),
]

_REQUIRED_SECTION_KEYS: frozenset = frozenset({
    "objetivo", "alcance", "responsabilidades",
    "procedimiento", "gestion_desviaciones", "registros",
})


# ---------------------------------------------------------------------------
# Error class
# ---------------------------------------------------------------------------

class DocxBuildError(Exception):
    """Raised when the DOCX file cannot be built due to a technical failure."""
    pass


# ---------------------------------------------------------------------------
# Internal helpers
# ---------------------------------------------------------------------------

def _clear_body_paragraphs(doc) -> None:
    """
    Remove all <w:p> elements from the document body, preserving:
      - <w:tbl>   — institutional metadata tables in the template
      - <w:sectPr> — section properties holding header/footer references

    This keeps the institutional chrome intact while making room for the
    generated content.
    """
    body = doc.element.body
    paragraphs_to_remove = [
        child for child in list(body.iterchildren())
        if child.tag.split("}")[-1] == "p"
    ]
    for elem in paragraphs_to_remove:
        body.remove(elem)


def _add_section_heading(doc, text: str) -> None:
    """Add a Heading 1 paragraph using the template's Heading 1 style."""
    try:
        doc.add_heading(text, level=1)
    except KeyError:
        # Template has no Heading 1 style — fall back to Normal + bold
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.bold = True


def _add_section_content(doc, text: str) -> None:
    """
    Add body content paragraphs.
    Each non-blank line in text becomes a separate paragraph, which preserves
    bullet lists and sub-steps that Claude typically formats as separate lines.
    """
    for line in text.split("\n"):
        if line.strip():
            doc.add_paragraph(line.strip())


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------

def build_docx_from_template(
    template_path: str,
    output_path: str,
    document_data: dict,
    document_text: str | None = None,
) -> dict:
    """
    Render the approved POE into an institutional .docx using the template.

    Parameters
    ----------
    template_path   Path to the institutional template (.docx).
    output_path     Destination path for the generated final.docx.
    document_data   Structured data from document.json
                    (must contain "sections": dict[str, str]).
    document_text   Raw text of the final approved document (final.txt).
                    Used as fallback when document_data has no sections.

    Returns
    -------
    {
      "ok":               True,
      "output_path":      str  — resolved path of the written file,
      "sections_written": list[str] — canonical keys rendered,
      "warnings":         list[str] — non-fatal issues,
    }

    Raises
    ------
    DocxBuildError  Template missing, unreadable, or output unwritable.
    """
    if not _DOCX_AVAILABLE:  # pragma: no cover
        raise DocxBuildError(
            "python-docx no está instalado. Instala con: pip install python-docx"
        )

    template_path_obj = Path(template_path)
    if not template_path_obj.exists():
        raise DocxBuildError(
            f"Plantilla DOCX no encontrada: {template_path!r}. "
            "Proporcione una plantilla institucional en la ruta configurada."
        )

    try:
        doc = Document(str(template_path_obj))
    except Exception as exc:
        raise DocxBuildError(
            f"No se pudo abrir la plantilla '{template_path}': {exc}"
        ) from exc

    warnings: list[str] = []
    sections_written: list[str] = []

    sections: dict[str, str] = document_data.get("sections", {})

    # Clear editable body paragraphs while preserving tables and header/footer.
    _clear_body_paragraphs(doc)

    if sections:
        for section_key, heading_label in _SECTION_RENDER_ORDER:
            content = sections.get(section_key, "").strip()
            if content:
                _add_section_heading(doc, heading_label)
                _add_section_content(doc, content)
                sections_written.append(section_key)
            elif section_key in _REQUIRED_SECTION_KEYS:
                warnings.append(
                    f"Sección requerida '{section_key}' ausente en document_data. "
                    "El documento puede estar incompleto."
                )
    else:
        # No structured sections — fall back to raw document_text
        if document_text and document_text.strip():
            warnings.append(
                "document_data no contiene secciones estructuradas. "
                "Se usó document_text como contenido completo (fallback sin estructura)."
            )
            _add_section_content(doc, document_text)
        else:
            warnings.append(
                "Ni document_data ni document_text contienen contenido. "
                "El DOCX generado estará vacío."
            )

    # Write output
    output_path_obj = Path(output_path)
    output_path_obj.parent.mkdir(parents=True, exist_ok=True)

    try:
        doc.save(str(output_path_obj))
    except Exception as exc:
        raise DocxBuildError(
            f"No se pudo guardar el DOCX en '{output_path}': {exc}"
        ) from exc

    return {
        "ok": True,
        "output_path": str(output_path_obj),
        "sections_written": sections_written,
        "warnings": warnings,
    }
