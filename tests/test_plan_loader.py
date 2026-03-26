"""Sprint 2B — tests for plan_loader.py

All tests are offline (no API calls, no real .xlsx files shipped with the repo).
Test xlsx files are created in tmp_path using openpyxl.
"""
import pytest
import openpyxl

from src.plan_loader import load_plan_xlsx, PlanLoadError


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

_REQUIRED_HEADERS = [
    "equipo_id", "familia_equipo", "actividad_id",
    "actividad_nombre", "frecuencia", "responsable", "registro_codigo",
]

def _make_xlsx(tmp_path, rows, headers=None, sheet_name="Plan_Mantenimiento"):
    """Create a minimal .xlsx in tmp_path with given headers and data rows."""
    if headers is None:
        headers = _REQUIRED_HEADERS
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.append(headers)
    for row in rows:
        ws.append(row)
    path = str(tmp_path / "plan.xlsx")
    wb.save(path)
    return path


def _one_row(
    equipo_id="EQ-001",
    familia="FAM:GAS",
    act_id="ACT-001",
    nombre="Inspección visual",
    frecuencia="mensual",
    responsable="Técnico de mantenimiento",
    registro="REG-001",
):
    return [equipo_id, familia, act_id, nombre, frecuencia, responsable, registro]


# ---------------------------------------------------------------------------
# Basic success path
# ---------------------------------------------------------------------------

def test_load_valid_plan_returns_dict(tmp_path):
    path = _make_xlsx(tmp_path, [_one_row()])
    result = load_plan_xlsx(path)
    assert isinstance(result, dict)


def test_load_valid_plan_fields(tmp_path):
    path = _make_xlsx(tmp_path, [_one_row()])
    result = load_plan_xlsx(path)
    assert result["familia_equipo"] == "FAM:GAS"
    assert result["sheet_name"] == "Plan_Mantenimiento"
    assert len(result["actividades"]) == 1
    assert result["actividades"][0]["actividad_id"] == "ACT-001"
    assert result["actividades"][0]["frecuencia"] == "mensual"
    assert "source_path" in result
    assert "loaded_at" in result


def test_empty_path_returns_none():
    assert load_plan_xlsx("") is None
    assert load_plan_xlsx(None) is None


def test_nombre_norm_is_normalized(tmp_path):
    path = _make_xlsx(tmp_path, [_one_row(nombre="Inspección visual de válvulas")])
    result = load_plan_xlsx(path)
    nombre_norm = result["actividades"][0]["nombre_norm"]
    # accents removed, lowercased
    assert "inspeccion" in nombre_norm
    assert "valvulas" in nombre_norm


def test_responsable_norm_is_normalized(tmp_path):
    path = _make_xlsx(tmp_path, [_one_row(responsable="Técnico de Mantención")])
    result = load_plan_xlsx(path)
    norm = result["actividades"][0]["responsable_norm"]
    assert "tecnico" in norm
    assert "mantencion" in norm


def test_frecuencia_canonical_mensual(tmp_path):
    for raw in ["mensual", "monthly", "cada mes", "1/mes"]:
        path = _make_xlsx(tmp_path, [_one_row(frecuencia=raw)])
        r = load_plan_xlsx(path)
        assert r["actividades"][0]["frecuencia"] == "mensual", f"failed for raw={raw!r}"


def test_frecuencia_canonical_trimestral(tmp_path):
    path = _make_xlsx(tmp_path, [_one_row(frecuencia="quarterly")])
    r = load_plan_xlsx(path)
    assert r["actividades"][0]["frecuencia"] == "trimestral"


def test_frecuencia_raw_preserved(tmp_path):
    path = _make_xlsx(tmp_path, [_one_row(frecuencia="Cada Mes")])
    r = load_plan_xlsx(path)
    assert r["actividades"][0]["frecuencia_raw"] == "Cada Mes"


def test_derived_views_are_sorted_lists(tmp_path):
    rows = [
        _one_row(act_id="ACT-001", responsable="Técnico B", registro="REG-002"),
        _one_row(act_id="ACT-002", responsable="Técnico A", registro="REG-001"),
    ]
    path = _make_xlsx(tmp_path, rows)
    r = load_plan_xlsx(path)
    assert r["responsables"] == sorted(r["responsables"])
    assert r["registros"] == sorted(r["registros"])
    assert r["frecuencias"] == sorted(r["frecuencias"])
    assert r["nombres_norm"] == sorted(r["nombres_norm"])


def test_familia_equipo_uppercased(tmp_path):
    path = _make_xlsx(tmp_path, [_one_row(familia="fam:gas")])
    r = load_plan_xlsx(path)
    assert r["familia_equipo"] == "FAM:GAS"


# ---------------------------------------------------------------------------
# Error cases
# ---------------------------------------------------------------------------

def test_file_not_found_raises_plan_load_error():
    with pytest.raises(PlanLoadError, match="no encontrado"):
        load_plan_xlsx("/tmp/nonexistent_plan_99999.xlsx")


def test_missing_required_column_raises(tmp_path):
    # Drop 'responsable' from headers
    headers = [h for h in _REQUIRED_HEADERS if h != "responsable"]
    rows = [[v for i, v in enumerate(_one_row()) if _REQUIRED_HEADERS[i] != "responsable"]]
    path = _make_xlsx(tmp_path, rows, headers=headers)
    with pytest.raises(PlanLoadError, match="responsable"):
        load_plan_xlsx(path)


def test_duplicate_actividad_id_raises(tmp_path):
    rows = [_one_row(act_id="ACT-DUP"), _one_row(act_id="ACT-DUP")]
    path = _make_xlsx(tmp_path, rows)
    with pytest.raises(PlanLoadError, match="duplicado"):
        load_plan_xlsx(path)


def test_multiple_familia_equipo_raises(tmp_path):
    rows = [
        _one_row(act_id="ACT-001", familia="FAM:GAS"),
        _one_row(act_id="ACT-002", familia="FAM:BIO"),
    ]
    path = _make_xlsx(tmp_path, rows)
    with pytest.raises(PlanLoadError, match="múltiples familias"):
        load_plan_xlsx(path)


def test_empty_sheet_raises(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Plan_Mantenimiento"
    path = str(tmp_path / "empty.xlsx")
    wb.save(path)
    with pytest.raises(PlanLoadError, match="vacía"):
        load_plan_xlsx(path)


def test_header_only_no_data_rows_raises(tmp_path):
    path = _make_xlsx(tmp_path, [])  # header row only, no data
    with pytest.raises(PlanLoadError, match="filas de datos válidas"):
        load_plan_xlsx(path)


def test_empty_required_field_in_row_raises(tmp_path):
    path = _make_xlsx(tmp_path, [_one_row(nombre="")])
    with pytest.raises(PlanLoadError, match="actividad_nombre"):
        load_plan_xlsx(path)


def test_blank_rows_are_skipped(tmp_path):
    """Blank separator rows between data rows should not cause an error."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Plan_Mantenimiento"
    ws.append(_REQUIRED_HEADERS)
    ws.append(list(_one_row(act_id="ACT-001")))
    ws.append([None] * len(_REQUIRED_HEADERS))  # blank separator
    ws.append(list(_one_row(act_id="ACT-002")))
    path = str(tmp_path / "plan_blanks.xlsx")
    wb.save(path)
    r = load_plan_xlsx(path)
    assert len(r["actividades"]) == 2


def test_column_header_aliases_accepted(tmp_path):
    """Alternative column header names should be accepted."""
    alt_headers = [
        "equipment_id", "familia", "activity_id",
        "activity_name", "frequency", "responsible", "formulario",
    ]
    path = _make_xlsx(tmp_path, [_one_row()], headers=alt_headers)
    r = load_plan_xlsx(path)
    assert r["familia_equipo"] == "FAM:GAS"
    assert r["actividades"][0]["actividad_id"] == "ACT-001"


def test_preferred_sheet_selected_over_first(tmp_path):
    """If 'Plan_Mantenimiento' sheet exists, it should be selected over the first sheet."""
    wb = openpyxl.Workbook()
    # First sheet (active by default) — wrong data
    ws_first = wb.active
    ws_first.title = "Datos_Generales"
    ws_first.append(["irrelevant", "columns"])
    # Second sheet — the real plan
    ws_plan = wb.create_sheet("Plan_Mantenimiento")
    ws_plan.append(_REQUIRED_HEADERS)
    ws_plan.append(list(_one_row()))
    path = str(tmp_path / "multi_sheet.xlsx")
    wb.save(path)
    r = load_plan_xlsx(path)
    assert r["sheet_name"] == "Plan_Mantenimiento"
    assert len(r["actividades"]) == 1
